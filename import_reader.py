"""通用账单读取：保留单元格值及来源元信息，不识别平台、不作行处理结论。

输出 BillFile → RawRow → RawCell；CSV 值为文本，XLSX 保留 openpyxl 返回的类型及格式。
行号从 1 开始，空行/说明行不自行跳过；字段映射、金额方向和导入决定由后续流程处理。
只读来源文件，不写账本或偏好；缺文件、解析错误等向调用方传播，不能返回伪造的空账单。
"""

import csv
from dataclasses import dataclass
import io
from pathlib import Path


@dataclass(frozen=True)
class RawCell:
    """读取值、类型标记及 Excel 数字格式；格式是判断线索，不是格式化后的显示文本。"""
    value: object
    kind: str
    number_format: str = ""


@dataclass(frozen=True)
class RawRow:
    """携原始起始行号的一行；CSV 跨行字段按起始物理行编号，XLSX 按工作表行号。"""
    row_number: int
    cells: tuple[RawCell, ...]

    @property
    def values(self) -> tuple[object, ...]:
        """仅提取值而不转字符串；需要区分公式/日期等类型时仍须读取 cells 元信息。"""
        return tuple(cell.value for cell in self.cells)


@dataclass(frozen=True)
class BillFile:
    """单个文件及所选明细的读取结果；encoding 仅 CSV 有值，sheet_name 仅 XLSX 有值。"""
    file_name: str
    rows: tuple[RawRow, ...]
    encoding: str | None = None
    sheet_name: str | None = None


def _read_csv(path: Path) -> BillFile:
    # 先完整解码再解析，编码回退只处理解码失败，不能把非法 CSV 格式误当成另一种编码。
    content = path.read_bytes()
    # UTF-8 优先且去 BOM，GB18030 兼容常见中文导出；不替换坏字节，以免金额/单号被静默改写。
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue  # 只试下一种支持编码，全部失败交给调用方提示，不猜更多编码。
    else:
        raise ValueError("CSV 编码无法读取，需要 UTF-8 或 GB18030 文件。")
    # newline="" 交给 csv 处理引号内换行；strict 保留格式错误，不手工按逗号或换行拆记录。
    reader = csv.reader(io.StringIO(text, newline=""), strict=True)
    rows = []
    next_line = 1
    for values in reader:
        # 引号内换行会占多条物理行；记录起始行号，不能以交易序号冒充原始行号。
        rows.append(RawRow(next_line, tuple(RawCell(value, "text") for value in values)))
        next_line = reader.line_num + 1
    return BillFile(path.name, tuple(rows), encoding=encoding)


def _read_xlsx(path: Path, sheet_name: str | None) -> BillFile:
    from openpyxl import load_workbook  # 延迟依赖，普通记账启动不读取 XLSX 或初始化解析器。

    # data_only=False 保留公式表达式供后续判定，不能把缓存计算值冒充原始确认值。
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        if sheet_name is None:
            # 单 sheet 才能隐式选取；多 sheet 必须由调用方明确选择，不能默认取首张漏掉来源。
            if len(workbook.sheetnames) != 1:
                raise ValueError("此文件有多个工作表，请明确指定一个，不合并工作表。")
            sheet_name = workbook.sheetnames[0]
        if sheet_name not in workbook.sheetnames:
            raise ValueError("指定工作表不存在。")
        sheet = workbook[sheet_name]
        # 外部导出器可能把范围声明写小；按实际 XML 读取，不能静默漏掉范围之外的交易。
        sheet.reset_dimensions()
        # read_only 保留日期类型和单元格格式；不执行公式、不把说明区当平台判断依据。
        rows = tuple(
            RawRow(index, tuple(
                RawCell(cell.value, cell.data_type, cell.number_format or "")
                for cell in cells
            ))
            for index, cells in enumerate(sheet.iter_rows(), start=1)
        )
        # 不依赖不可靠的声明宽度；按实际最大列数补齐空单元格，保留列位置和中间空行。
        width = max((len(row.cells) for row in rows), default=0)
        rows = tuple(
            RawRow(row.row_number, row.cells + (RawCell(None, "n"),) * (width - len(row.cells)))
            for row in rows
        )
        return BillFile(path.name, rows, sheet_name=sheet_name)
    finally:
        workbook.close()  # 关闭后取消／重读都不会遗留被 Windows 占用的文件句柄。


def read_bill_file(path: Path, *, sheet_name: str | None = None) -> BillFile:
    """按扩展名只读 CSV/XLSX，返回完整读取结果；不检查交易语义或写入数据库。"""
    path = Path(path)
    if path.suffix.casefold() == ".csv":
        if sheet_name is not None:
            raise ValueError("CSV 没有工作表。")
        return _read_csv(path)
    if path.suffix.casefold() == ".xlsx":
        return _read_xlsx(path, sheet_name)
    raise ValueError("本次只支持 CSV 和 XLSX 账单文件。")
